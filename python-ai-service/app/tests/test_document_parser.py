# 测试解析docx文件返回纯文本
from docx import Document

from app.services.review_engine import read_unstructured_text


def test_parse_docx_returns_plain_text(tmp_path):
    file_path = tmp_path / "取水许可申请书.docx"

    document = Document()
    document.add_heading("取水许可申请书", level=1)
    document.add_paragraph("申请人：测试用户甲")
    document.add_paragraph("项目名称：农业灌溉取水项目")
    document.add_paragraph("取水地点：测试省测试市测试区")
    document.save(file_path)

    text = read_unstructured_text(file_path)

    assert "取水许可申请书" in text
    assert "申请人：测试用户甲" in text
    assert "农业灌溉取水项目" in text
    assert "取水地点" in text
