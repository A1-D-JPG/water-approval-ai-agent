import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app


client = TestClient(app)


@pytest.mark.integration
def test_review_api_processes_a_real_docx(tmp_path, monkeypatch):
    """临时生成 DOCX，并通过 /review 跑通完整审核流程。"""
    monkeypatch.setenv("OPENAI_API_KEY", "")


    file_path = tmp_path / "取水许可申请书.docx"

    document = Document()
    document.add_heading("取水许可申请书", level=1)
    document.add_paragraph("申请人姓名：测试用户甲")
    document.add_paragraph("身份证号：TEST-ID-0001")
    document.add_paragraph("联系方式：TEST-PHONE-0001")
    document.add_paragraph("项目名称：农业灌溉取水项目")
    document.add_paragraph("取水地点：测试省测试市测试区")
    document.add_paragraph("取水用途：农业灌溉")
    document.add_paragraph("行业类别：011谷物种植")
    document.save(file_path)

    response = client.post(
        "/review",
        headers={"X-Request-Id": "complete-review-test"},
        json={
            "application_id": 1001,
            "applicant_name": "测试用户甲",
            "id_number": "TEST-ID-0001",
            "project_name": "农业灌溉取水项目",
            "water_location": "测试省测试市测试区",
            "water_use": "农业灌溉",
            "industry_category": "011谷物种植",
            "contact_phone": "TEST-PHONE-0001",
            "credit_code": "",
            "file_paths": [str(file_path)],
        },
    )

    assert response.status_code == 200

    body = response.json()

    assert body["application_id"] == 1001
    assert body["applicant_name"] == "测试用户甲"
    assert body["ai_status"] in {
        "APPROVED",
        "REJECTED",
        "NEED_MANUAL_REVIEW",
    }

    assert isinstance(body["issues"], list)
    assert isinstance(body["suggestions"], list)
    assert isinstance(body["knowledge_hits"], list)
    assert body["raw_length"] > 0

    assert len(body["files"]) == 1
    assert body["files"][0]["name"] == "取水许可申请书.docx"

    assert body["agent_result"]["mode"] == "deterministic_rule_agent"
    assert body["agent_result"]["decision"] == body["ai_status"]
    assert body["agent_result"]["analysis"]

    assert response.headers["X-Request-Id"] == "complete-review-test"
